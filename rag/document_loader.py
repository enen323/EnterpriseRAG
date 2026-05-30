import logging
from pathlib import Path
from typing import List

from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_core.documents import Document as LCDocument

from core.config import settings

logger = logging.getLogger(__name__)


def load_document(file_path: str) -> List[LCDocument]:
    """Load a document file and return LangChain Document objects."""
    path = Path(file_path)
    suffix = path.suffix.lower()

    if suffix == ".txt":
        return _load_txt(path)
    elif suffix == ".md":
        return _load_md(path)
    elif suffix == ".pdf":
        return _load_pdf(path)
    elif suffix == ".docx":
        return _load_docx(path)
    else:
        raise ValueError(f"Unsupported file type: {suffix}")


def _load_txt(path: Path) -> List[LCDocument]:
    text = path.read_text(encoding="utf-8")
    return [LCDocument(page_content=text, metadata={"source": path.name})]


def _load_md(path: Path) -> List[LCDocument]:
    text = path.read_text(encoding="utf-8")
    return [LCDocument(page_content=text, metadata={"source": path.name})]


def _load_pdf(path: Path) -> List[LCDocument]:
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    docs = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text()
        if text.strip():
            docs.append(LCDocument(page_content=text, metadata={"source": path.name, "page": i + 1}))
    return docs


def _load_docx(path: Path) -> List[LCDocument]:
    from docx import Document as DocxDocument
    doc = DocxDocument(str(path))
    text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
    return [LCDocument(page_content=text, metadata={"source": path.name})]


def split_documents(docs: List[LCDocument]) -> List[LCDocument]:
    """Split documents into chunks."""
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=settings.CHUNK_SIZE,
        chunk_overlap=settings.CHUNK_OVERLAP,
        separators=["\n\n", "\n", "。", ".", " ", ""],
    )
    return splitter.split_documents(docs)
