import uuid
import logging
import tempfile
import os
import shutil
import zipfile
import io
from pathlib import Path
from typing import Annotated, List

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, status
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from core.database import get_db
from core.models import User, Document, DocumentStatus
from core.schemas import DocumentOut, PreviewOut
from core.config import settings
from api.deps import get_current_user
from rag.document_loader import load_document, split_documents
from rag.vector_store import add_documents, delete_document_chunks

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api/documents", tags=["documents"])

ALLOWED_EXTENSIONS = {".pdf", ".md", ".txt", ".docx"}


@router.get("", response_model=List[DocumentOut])
async def list_documents(
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    result = await db.execute(
        select(Document).where(Document.user_id == current_user.id).order_by(Document.created_at.desc())
    )
    return result.scalars().all()


@router.post("/upload", response_model=DocumentOut, status_code=status.HTTP_201_CREATED)
async def upload_document(
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
    file: UploadFile = File(...),
):
    ext = os.path.splitext(file.filename or "")[1].lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=400, detail=f"Unsupported file type: {ext}. Allowed: {ALLOWED_EXTENSIONS}")

    doc_id = uuid.uuid4()
    doc_record = Document(
        id=doc_id,
        user_id=current_user.id,
        filename=file.filename or "unknown",
        file_type=ext,
        status=DocumentStatus.PROCESSING,
    )
    db.add(doc_record)
    await db.commit()

    tmp_path = None
    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as tmp:
            tmp.write(await file.read())
            tmp_path = tmp.name

        # Save a copy to persistent storage for preview
        storage_dir = Path(settings.STORAGE_DIR) / str(current_user.id)
        storage_dir.mkdir(parents=True, exist_ok=True)
        storage_path = storage_dir / f"{doc_id}{ext}"
        shutil.copy2(tmp_path, storage_path)
        doc_record.storage_path = str(storage_path)

        raw_docs = load_document(tmp_path)
        chunks = split_documents(raw_docs)
        chunk_count = add_documents(chunks, str(doc_id), user_id=str(current_user.id))

        doc_record.status = DocumentStatus.READY
        doc_record.chunk_count = chunk_count
        await db.commit()
    except Exception as e:
        doc_record.status = DocumentStatus.FAILED
        await db.commit()
        logger.error(f"Document ingestion failed: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=f"Ingestion failed: {str(e)}")
    finally:
        if tmp_path and os.path.exists(tmp_path):
            os.unlink(tmp_path)

    await db.refresh(doc_record)
    return doc_record


@router.get("/{doc_id}/preview", response_model=PreviewOut)
async def preview_document(
    doc_id: uuid.UUID,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    result = await db.execute(
        select(Document).where(Document.id == doc_id, Document.user_id == current_user.id)
    )
    doc = result.scalar_one_or_none()
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found")

    if doc.file_type in (".md", ".txt"):
        content = Path(doc.storage_path).read_text(encoding="utf-8") if doc.storage_path else "File not available"
    elif doc.file_type == ".docx":
        from docx import Document as DocxDocument
        d = DocxDocument(doc.storage_path) if doc.storage_path else None
        content = "\n".join(p.text for p in d.paragraphs if p.text.strip()) if d else "File not available"
    else:
        raise HTTPException(status_code=400, detail="Preview not available for this file type")

    return PreviewOut(content=content, filename=doc.filename, file_type=doc.file_type)


@router.post("/upload/batch", response_model=List[DocumentOut], status_code=status.HTTP_201_CREATED)
async def upload_documents_batch(
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
    files: List[UploadFile] = File(...),
):
    results = []
    for file in files:
        try:
            ext = os.path.splitext(file.filename or "")[1].lower()
            if ext not in ALLOWED_EXTENSIONS and ext != ".zip":
                continue

            if ext == ".zip":
                zip_results = await _process_zip(file, current_user, db)
                results.extend(zip_results)
                continue

            # Single file (reuse logic from upload_document)
            doc_id = uuid.uuid4()
            doc_record = Document(
                id=doc_id, user_id=current_user.id, filename=file.filename or "unknown",
                file_type=ext, status=DocumentStatus.PROCESSING,
            )
            db.add(doc_record)
            await db.flush()

            tmp_path = None
            try:
                with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as tmp:
                    tmp.write(await file.read())
                    tmp_path = tmp.name

                # Save to persistent storage for preview
                storage_dir = Path(settings.STORAGE_DIR) / str(current_user.id)
                storage_dir.mkdir(parents=True, exist_ok=True)
                storage_path = storage_dir / f"{doc_id}{ext}"
                shutil.copy2(tmp_path, storage_path)
                doc_record.storage_path = str(storage_path)

                raw_docs = load_document(tmp_path)
                chunks = split_documents(raw_docs)
                chunk_count = add_documents(chunks, str(doc_id), user_id=str(current_user.id))

                doc_record.status = DocumentStatus.READY
                doc_record.chunk_count = chunk_count
            except Exception as e:
                doc_record.status = DocumentStatus.FAILED
                logger.error(f"Batch doc failed: {file.filename}: {e}")
            finally:
                if tmp_path and os.path.exists(tmp_path):
                    os.unlink(tmp_path)

            await db.refresh(doc_record)
            results.append(doc_record)
        except Exception as e:
            logger.error(f"Batch upload file error: {file.filename}: {e}")

    await db.commit()
    return results


async def _process_zip(file: UploadFile, current_user: User, db: AsyncSession) -> list[Document]:
    results = []
    content = await file.read()
    with zipfile.ZipFile(io.BytesIO(content)) as zf:
        for name in zf.namelist():
            ext = os.path.splitext(name)[1].lower()
            if ext not in ALLOWED_EXTENSIONS:
                continue

            doc_id = uuid.uuid4()
            doc_record = Document(
                id=doc_id, user_id=current_user.id, filename=name,
                file_type=ext, status=DocumentStatus.PROCESSING,
            )
            db.add(doc_record)
            await db.flush()

            tmp_path = None
            try:
                with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as tmp:
                    tmp.write(zf.read(name))
                    tmp_path = tmp.name

                # Save to persistent storage for preview
                storage_dir = Path(settings.STORAGE_DIR) / str(current_user.id)
                storage_dir.mkdir(parents=True, exist_ok=True)
                storage_path = storage_dir / f"{doc_id}{ext}"
                shutil.copy2(tmp_path, storage_path)
                doc_record.storage_path = str(storage_path)

                raw_docs = load_document(tmp_path)
                chunks = split_documents(raw_docs)
                chunk_count = add_documents(chunks, str(doc_id), user_id=str(current_user.id))

                doc_record.status = DocumentStatus.READY
                doc_record.chunk_count = chunk_count
            except Exception as e:
                doc_record.status = DocumentStatus.FAILED
                logger.error(f"ZIP entry failed: {name}: {e}")
            finally:
                if tmp_path and os.path.exists(tmp_path):
                    os.unlink(tmp_path)

            await db.refresh(doc_record)
            results.append(doc_record)
    return results


@router.delete("/{doc_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_document(
    doc_id: uuid.UUID,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    result = await db.execute(
        select(Document).where(Document.id == doc_id, Document.user_id == current_user.id)
    )
    doc = result.scalar_one_or_none()
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found")

    delete_document_chunks(str(doc_id))

    await db.delete(doc)
    await db.commit()


@router.get("/{doc_id}/status", response_model=DocumentOut)
async def get_document_status(
    doc_id: uuid.UUID,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    result = await db.execute(
        select(Document).where(Document.id == doc_id, Document.user_id == current_user.id)
    )
    doc = result.scalar_one_or_none()
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found")
    return doc
